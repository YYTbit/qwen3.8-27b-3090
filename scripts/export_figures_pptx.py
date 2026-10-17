import json
import statistics
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt


repository_root = Path(__file__).resolve().parents[1]
figure_data = json.loads((repository_root / 'docs/figures/figure-data.json').read_text())
presentation = Presentation()
presentation.slide_width = Inches(13.4)
presentation.slide_height = Inches(7)
presentation.core_properties.author = 'YYTbit'
presentation.core_properties.last_modified_by = 'YYTbit'
presentation.core_properties.title = 'Qwen3.8-27B Dual-GPU Performance Evaluation'


def rgb(value):
    return RGBColor.from_string(value)


def add_text(slide, x, y, w, h, value, font_size=13, bold=False, color='334155'):
    text_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    z = text_box.text_frame
    z.clear()
    z.word_wrap = True
    z.margin_left = z.margin_right = z.margin_top = z.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = z.paragraphs[0] if i == 0 else z.add_paragraph()
        p.text = line
        p.font.name = 'Arial'
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(4)
    return text_box


def add_shape(slide, x, y, w, h, fill_color, border_color=None, shape_type=MSO_SHAPE.RECTANGLE):
    z = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    z.fill.solid()
    z.fill.fore_color.rgb = rgb(fill_color)
    if border_color:
        z.line.color.rgb = rgb(border_color)
        z.line.width = Pt(1)
    else:
        z.line.fill.background()
    return z


def add_line(slide, x, y, a, b, color='CBD5E1', line_width=1):
    z = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(a), Inches(b))
    z.line.color.rgb = rgb(color)
    z.line.width = Pt(line_width)


def add_slide(title, subtitle):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(slide, 0.4, 0.22, 12.5, 0.5, title, 25, True, '0F172A')
    add_text(slide, 0.4, 0.86, 12.4, 0.5, subtitle, 12, False, '475569')
    return slide


slide = add_slide('Dual-GPU Evaluation Overview',
    'Engine compatibility, context scaling, speculative decoding and deployment validation')
positions = [(0.02 + i * 0.25, 0.51) for i in range(4)] + [(0.02 + i * 0.25, 0.15) for i in [3, 2, 1, 0]]
for i, ((x, y), (tag, title, body)) in enumerate(zip(positions, figure_data['decision_cards'])):
    a, b = 13.4 * x, 7 * (1 - y - 0.295)
    fill_color, border_color = ('F8FAFC', 'CBD5E1')
    if i == 5:
        fill_color, border_color = 'ECFDF5', '059669'
    elif i == 6:
        fill_color, border_color = 'F0FDFA', '5EEAD4'
    elif i == 7:
        fill_color, border_color = 'EFF6FF', '93C5FD'
    add_shape(slide, a, b, 13.4 * 0.215, 7 * 0.295, fill_color, border_color, MSO_SHAPE.ROUNDED_RECTANGLE)
    add_text(slide, a + 0.13, b + 0.15, 2.62, 0.2, tag, 9.5, True, '166534' if i == 5 else '64748B')
    add_text(slide, a + 0.13, b + 0.5, 2.62, 0.43, title, 13.2, True, '0F172A')
    add_text(slide, a + 0.13, b + 1.01, 2.62, 0.93, body, 11.6)
    if i < 3:
        add_shape(slide, a + 2.98, b + 0.97, 0.24, 0.11, '64748B', shape_type=MSO_SHAPE.RIGHT_ARROW)
    elif i == 3:
        add_shape(slide, a + 1.37, b + 2.13, 0.11, 0.23, '64748B', shape_type=MSO_SHAPE.DOWN_ARROW)
    elif i < 7:
        add_shape(slide, a - 0.33, b + 0.97, 0.24, 0.11, '64748B', shape_type=MSO_SHAPE.LEFT_ARROW)
add_text(slide, 0.4, 6.48, 12.4, 0.36, 'Campaigns overlap. Results retain their configurations. Headline: 18,649 actual input tokens, warm prefix, rates include reasoning.', 10.6)

slide = add_slide('Long-Context Prefill Scaling',
    'vLLM 0.23 W4A8 PP2 | Median completed-batch prefill rate by processed-token range')
x0, y0, w, h = 1.05, 5.35, 11.85, 3.45
for v in [0, 500, 1000, 1500, 2000, 2500]:
    y = y0 - h * v / 2900
    add_line(slide, x0, y, x0 + w, y, 'E2E8F0', 0.7)
    add_text(slide, 0.38, y - 0.09, 0.57, 0.2, str(v), 10.5)
for i, v in enumerate(figure_data['prefill']):
    value = v['median_batch_tokens_per_s']
    x = x0 + (i + 0.5) * w / 7
    bar_height = h * value / 2900
    add_shape(slide, x - 0.51, y0 - bar_height, 1.02, bar_height, '2563EB' if i == 0 else '0F766E' if i == 6 else '94A3B8')
    add_text(slide, x - 0.3, y0 - bar_height - 0.31, 1.0, 0.25, '{:,.0f}'.format(value), 14, True, '0F172A')
    a, b = v['token_range']
    s = '{}-{}k'.format(a // 1000, round(b / 1000, 1) if b % 1000 else b // 1000)
    add_text(slide, x - 0.6, y0 + 0.10, 1.4, 0.23, s, 10.5)
add_text(slide, 0.45, 1.47, 5, 0.23, 'Input tokens / second', 11.5)
add_text(slide, 5.2, 5.86, 4.5, 0.25, 'Processed prompt tokens', 12)
add_text(slide, 0.45, 6.31, 12.4, 0.45, 'Actual input: 520,761 tokens. First batch excluded. The 3,437 tok/s early window uses a different metric. Source: S038.', 10.6)

slide = add_slide('SGLang Parallelism Evaluation',
    'SGLang 0.5.19 + MTP | 5,224 input / 4,096 output tokens | One observation per layout')
x0, w = 1.35, 10.9
for v in range(0, 61, 10):
    x = x0 + w * v / 60
    add_line(slide, x, 1.85, x, 5.15, 'E2E8F0', 0.7)
    add_text(slide, x - 0.08, 5.3, 0.5, 0.23, str(v), 11)
for i, v in enumerate(figure_data['sglang_layouts']):
    y = 2.35 + i * 1.6
    rate = v['first_api_tokens_per_s']
    add_shape(slide, x0, y, w * rate / 60, 0.65, '2563EB' if i == 0 else '0F766E')
    add_text(slide, 0.58, y + 0.18, 0.7, 0.25, v['layout'], 14, True)
    add_text(slide, x0 + w * rate / 60 + 0.12, y + 0.19, 0.9, 0.3, '{:.2f}'.format(rate), 15, True, '0F172A')
add_text(slide, 3.9, 5.85, 7.0, 0.28, 'Output tokens / second, including reasoning', 12)
add_text(slide, 0.45, 6.33, 12.4, 0.45, 'PP2: AWQ W4A16, FP8 KV, historical backports. Complete TP2 launch recovery is unavailable. Source: S016.', 10.6)

slide = add_slide('Archive Restoration and Reconstruction',
    '18,649 input / 1,024 output tokens | Three observations per group | Rates include reasoning')
x0, y0, w, h = 1.05, 5.25, 11.75, 3.4
add_shape(slide, x0 + 3 * w / 4, y0 - h, w / 4, h, 'F1F5F9')
for v in range(120, 171, 10):
    y = y0 - h * (v - 120) / 50
    add_line(slide, x0, y, x0 + w, y, 'E2E8F0', 0.7)
    add_text(slide, 0.43, y - 0.1, 0.5, 0.24, str(v), 10.5)
for i, (label, value) in enumerate(figure_data['restoration_groups']):
    x = x0 + (i + 0.5) * w / 4
    color = '2563EB' if i == 0 else '0F766E' if i < 3 else '64748B'
    median_rate = statistics.median(value)
    y = y0 - h * (median_rate - 120) / 50
    add_line(slide, x - 0.5, y, x + 0.5, y, '0F172A', 1.8)
    add_text(slide, x + 0.58, y - 0.09, 0.8, 0.25, '{:.2f}'.format(median_rate), 12.6, True, '0F172A')
    for j, v in enumerate(value):
        a = x + (j - 1) * 0.34
        b = y0 - h * (v - 120) / 50
        add_shape(slide, a - 0.062, b - 0.062, 0.124, 0.124, color, 'FFFFFF', MSO_SHAPE.OVAL)
    add_text(slide, x - 0.85, 5.43, 2.0, 0.6, label, 12)
add_text(slide, 0.45, 1.43, 7.0, 0.26, 'Output tokens / second, including reasoning', 11.5)
add_text(slide, 0.45, 6.25, 12.4, 0.56, 'Dots: individual runs. Black segments: medians. First three groups: 17,920 prefix hits. Separate assembly: 16,128 hits, 10.5 GiB KV/card.', 10.6)

presentation.save(repository_root / 'docs/figures/dual-gpu-exploration.pptx')
print('Saved four editable slides with native shapes and text.')
